"""Excel rendering for normalized offline AD analysis reports."""

from __future__ import annotations

import json
import re
from typing import Any, Iterable, List, Mapping, Sequence

from .ad_excel_views import add_operator_views


def _excel_value(value: Any) -> Any:
    """Return an inert, legal Excel cell value."""
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, sort_keys=True, default=str)
    text = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", str(value))
    if len(text) > 32767:
        text = text[:32740] + "... [truncated]"
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _collect_paths(report: Mapping[str, Any]) -> List[Mapping[str, Any]]:
    """Collect unique path evidence from either report location."""
    paths: List[Mapping[str, Any]] = []
    seen = set()
    candidates = list(report.get("owned_analysis", {}).get("paths", []))
    candidates.extend(
        report.get("privilege_analysis", {}).get("dcsync_paths", [])
    )
    for finding in report.get("findings", []):
        if str(finding.get("id", "")).endswith(
            ("path_to_high_value", "path_to_domain_admin")
        ):
            candidates.extend(finding.get("evidence", []))
    for path in candidates:
        if not isinstance(path, dict):
            continue
        key = (
            path.get("source", {}).get("id"),
            path.get("target", {}).get("id"),
            tuple(path.get("edges", [])),
        )
        if key not in seen:
            seen.add(key)
            paths.append(path)
    return paths


class ADExcelFormatter:
    """Build a normalized operator workbook from an AD analysis report."""

    def format(
        self,
        report: Mapping[str, Any],
        mapped_findings: Sequence[Mapping[str, Any]] = (),
    ) -> Any:
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.remove(workbook.active)
        mapping = {
            finding_id: mapped["internal_vulnerability_id"]
            for mapped in mapped_findings
            for finding_id in mapped["finding_ids"]
        }

        self._add_sheet(
            workbook,
            "Summary",
            ["Metric", "Value"],
            self._summary_rows(report, mapped_findings),
        )
        add_operator_views(
            self._add_sheet, workbook, report, mapping, _collect_paths(report)
        )

        self._style(workbook)
        return workbook

    @staticmethod
    def _add_sheet(
        workbook: Any,
        name: str,
        headers: Sequence[str],
        rows: Iterable[Sequence[Any]],
    ) -> Any:
        sheet = workbook.create_sheet(name)
        sheet.append(list(headers))
        for row in rows:
            sheet.append([_excel_value(value) for value in row])
        return sheet

    @staticmethod
    def _summary_rows(
        report: Mapping[str, Any], mapped: Sequence[Mapping[str, Any]]
    ) -> List[Sequence[Any]]:
        summary = report.get("summary", {})
        engine = report.get("engine", {})
        source = report.get("source", {})
        owned = report.get("owned_analysis", {})
        privilege = report.get("privilege_analysis", {})
        operator = report.get("operator_analysis", {})
        choke_points = report.get("path_analysis", {}).get("choke_points", [])
        return [
            ("Source", source.get("path", "")),
            ("Path backend", engine.get("path_backend", "")),
            ("Nodes", engine.get("node_count", 0)),
            ("Edges", engine.get("edge_count", 0)),
            ("Findings", summary.get("total", 0)),
            ("Critical", summary.get("critical", 0)),
            ("High", summary.get("high", 0)),
            ("Medium", summary.get("medium", 0)),
            ("Low", summary.get("low", 0)),
            ("Info", summary.get("info", 0)),
            ("Mapped internal findings", len(mapped)),
            ("Administrative memberships", len(privilege.get("memberships", []))),
            ("Administrative permissions", len(privilege.get("permissions", []))),
            ("DCSync paths", len(privilege.get("dcsync_paths", []))),
            ("Risk-filtered fleet access", len(operator.get("fleet_access", []))),
            ("AD CS objects", len(operator.get("adcs", {}).get("objects", []))),
            ("Path choke points", len(choke_points)),
            ("Owned principals", len(owned.get("resolved", []))),
            ("Owned paths", len(owned.get("paths", []))),
            ("Unresolved owned identities", len(owned.get("unresolved", []))),
            ("Ambiguous owned identities", len(owned.get("ambiguous", []))),
        ]

    @staticmethod
    def _style(workbook: Any) -> None:
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        alignment = Alignment(vertical="top", wrap_text=True)
        header_font = Font(color="FFFFFF", bold=True)
        severity_font = Font(color="FFFFFF", bold=True)
        header_fill = PatternFill("solid", fgColor="1F4E78")
        severity_fills = {
            "critical": PatternFill("solid", fgColor="C00000"),
            "high": PatternFill("solid", fgColor="FF0000"),
            "medium": PatternFill("solid", fgColor="FFC000"),
            "low": PatternFill("solid", fgColor="92D050"),
            "info": PatternFill("solid", fgColor="5B9BD5"),
        }
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            header = tuple(sheet[1])
            severity_column = next(
                (cell.column for cell in header if cell.value == "Severity"), None
            )
            widths = [0] * sheet.max_column
            for row_number, row in enumerate(sheet.iter_rows(), start=1):
                for cell in row:
                    cell.alignment = alignment
                    widths[cell.column - 1] = max(
                        widths[cell.column - 1], len(str(cell.value or ""))
                    )
                    if row_number == 1:
                        cell.font = header_font
                        cell.fill = header_fill
                if severity_column and row_number > 1:
                    cell = row[severity_column - 1]
                    fill = severity_fills.get(str(cell.value).lower())
                    if fill:
                        cell.fill = fill
                        cell.font = severity_font
            for column, content_width in enumerate(widths, start=1):
                sheet.column_dimensions[get_column_letter(column)].width = min(
                    60, max(10, content_width + 2)
                )
