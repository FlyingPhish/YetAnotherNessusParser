"""
yapp/core/excel_formatter.py
Excel report generator for consolidated vulnerability data and Nmap comparison
"""

import logging
from typing import Dict, Any, Optional
from openpyxl import Workbook

from .formatter import FormatterError

logger = logging.getLogger(__name__)


class ExcelFormatter:
    """
    Formats consolidated vulnerability data and Nmap comparison data into Excel workbooks.
    """
    
    def __init__(self):
        """Initialize Excel formatter"""
        self.max_sheet_name_length = 31  # Excel sheet name limit
    
    def format(self, consolidated_data: Dict[str, Any]) -> Optional[Workbook]:
        """
        Generate Excel workbook from consolidated vulnerability data.
        
        Args:
            consolidated_data: Consolidated vulnerability data with structure:
                {
                    'consolidated_vulnerabilities': {
                        'vuln_id': {
                            'consolidated_plugins': {...},
                            'affected_services': {...}
                        }
                    }
                }
        
        Returns:
            Workbook object ready to save, or None if no valid data
            
        Raises:
            FormatterError: If data validation or formatting fails
        """
        try:
            if not consolidated_data or not consolidated_data.get('consolidated_vulnerabilities'):
                logger.warning("No consolidated vulnerabilities found for Excel formatting")
                return None
            
            wb = Workbook()
            wb.remove(wb.active)  # Remove default sheet
            
            vulnerabilities = consolidated_data.get('consolidated_vulnerabilities', {})
            
            # Create sheet for each vulnerability
            for vuln_id, vuln_data in vulnerabilities.items():
                self._create_vulnerability_sheet(wb, vuln_id, vuln_data)
            
            logger.info(f"Successfully formatted {len(vulnerabilities)} vulnerabilities into Excel workbook")
            return wb
            
        except Exception as e:
            logger.error(f"Excel formatting failed: {str(e)}")
            raise FormatterError(f"Failed to format Excel: {str(e)}")
    
    def format_nmap_comparison(
        self,
        comparison_data: Dict[str, Any],
        first_filename: str,
        second_filename: str
    ) -> Optional[Workbook]:
        """
        Generate Excel workbook from Nmap comparison data.
        
        Args:
            comparison_data: Comparison data from NmapComparator with structure:
                {
                    'comparison_metadata': {...},
                    'comparison_data': [...],
                    'statistics': {...}
                }
            first_filename: Name of first scan file
            second_filename: Name of second scan file
            
        Returns:
            Workbook object ready to save, or None if no valid data
            
        Raises:
            FormatterError: If data validation or formatting fails
        """
        try:
            if not comparison_data or not comparison_data.get('comparison_data'):
                logger.warning("No comparison data found for Excel formatting")
                return None
            
            wb = Workbook()
            
            # Create comparison sheet
            self._create_comparison_sheet(
                wb,
                comparison_data['comparison_data'],
                first_filename,
                second_filename
            )
            
            logger.info(f"Successfully formatted Nmap comparison into Excel workbook")
            return wb
            
        except Exception as e:
            logger.error(f"Excel comparison formatting failed: {str(e)}")
            raise FormatterError(f"Failed to format Excel comparison: {str(e)}")
    
    def _create_comparison_sheet(
        self,
        wb: Workbook,
        comparison_data: list,
        first_filename: str,
        second_filename: str
    ) -> None:
        """
        Create comparison worksheet with port/service differences.
        
        Args:
            wb: Workbook to add sheet to
            comparison_data: List of comparison tuples
            first_filename: Name of first scan
            second_filename: Name of second scan
        """
        ws = wb.active
        ws.title = "Scan Comparison"
        
        # Add headers
        headers = [
            'IP Address',
            f'{first_filename} - Port/Protocol',
            f'{first_filename} - Service',
            f'{second_filename} - Port/Protocol',
            f'{second_filename} - Service',
            'Differences'
        ]
        ws.append(headers)
        
        # Add comparison data rows
        for row in comparison_data:
            ws.append(row)
        
        logger.debug(f"Created comparison sheet with {len(comparison_data)} rows")
    
    def _create_vulnerability_sheet(
        self, 
        wb: Workbook, 
        vuln_id: str, 
        vuln_data: Dict[str, Any]
    ) -> None:
        """
        Create a worksheet for a single vulnerability.
        
        Args:
            wb: Workbook to add sheet to
            vuln_id: Vulnerability identifier
            vuln_data: Vulnerability data including plugins and affected services
        """
        # Get plugin data
        plugins = vuln_data.get('consolidated_plugins', {})
        affected_services = vuln_data.get('affected_services', {})
        
        # Create safe sheet name (Excel 31 char limit)
        sheet_name = self._sanitize_sheet_name(vuln_id)
        ws = wb.create_sheet(title=sheet_name)
        
        # Build headers: FQDN, IP, Port, then plugin names
        headers = ['FQDN', 'IP', 'Port'] + list(plugins.values())
        ws.append(headers)
        
        # Add row for each affected service
        for service in affected_services.values():
            row = self._build_service_row(service, plugins)
            ws.append(row)
        
        logger.debug(f"Created sheet '{sheet_name}' with {len(affected_services)} services")
    
    def _build_service_row(
        self, 
        service: Dict[str, Any], 
        plugins: Dict[str, str]
    ) -> list:
        """
        Build a row for a service showing which plugins found it.
        
        Args:
            service: Service data with fqdn, ip, port, and issues_found
            plugins: Dict mapping plugin IDs to plugin names
            
        Returns:
            list: Row data [fqdn, ip, port, Yes/No for each plugin]
        """
        # Get set of plugin IDs found for this service
        found_ids = {issue['id'] for issue in service.get('issues_found', [])}
        
        # Build row: service info + Yes/No for each plugin
        row = [
            service.get('fqdn', ''),
            service.get('ip', ''),
            service.get('port', '')
        ]
        
        # Add Yes/No for each plugin
        row.extend(['Yes' if pid in found_ids else 'No' for pid in plugins.keys()])
        
        return row
    
    def _sanitize_sheet_name(self, name: str) -> str:
        """
        Sanitize sheet name to meet Excel requirements.
        
        Args:
            name: Original sheet name
            
        Returns:
            str: Sanitized sheet name (max 31 chars, no invalid chars)
        """
        # Remove invalid characters for Excel sheet names
        invalid_chars = ['\\', '/', '*', '?', ':', '[', ']']
        sanitized = name
        for char in invalid_chars:
            sanitized = sanitized.replace(char, '_')
        
        # Truncate to Excel's 31 character limit
        if len(sanitized) > self.max_sheet_name_length:
            sanitized = sanitized[:self.max_sheet_name_length]
        
        return sanitized